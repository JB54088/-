"""Mock verification for the administrator ABC Excel import workflow.

Run inside the backend container with:

    python -m unittest tests.test_excel_import -v
"""

import unittest
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from uuid import uuid4

from openpyxl import Workbook
from sqlalchemy import func, select

from app.database import SessionLocal
from app.consumption import calculate_customer_consumption_summary
from app.integrations.excel_import.parser import parse_excel
from app.integrations.excel_import.service import confirm_import, link_pending_external_customer, preview_import
from app.models import Customer, CustomerPoolMatchHistory, ExcelImportUnmatchedRecord, IntegrationLog, Order, OrderItem, User
from app.routers.excel_imports import unmatched_records


def workbook_bytes(external_customer_id: str, phone: str, transaction_id: str, mismatch: bool = False) -> bytes:
    workbook = Workbook()
    流水 = workbook.active
    流水.title = "流水"
    流水.append(["患者档案号", "患者姓名", "手机号", "收费单号", "收费时间", "支付方式", "实收", "基金支付总额", "非基金支付总额", "一级分类", "项目名称"])
    流水.append([external_customer_id, "Excel测试客户", phone, transaction_id, datetime(2026, 8, 26, 10, 0), "微信支付", 1500, 700, 200 if mismatch else 800, "问诊", "综合服务"])
    明细 = workbook.create_sheet("明细")
    明细.append(["患者档案号", "患者姓名", "手机号", "收费单号", "收费时间", "项目名称", "明细金额", "一级分类"])
    明细.append([external_customer_id, "Excel测试客户", phone, transaction_id, datetime(2026, 8, 26, 10, 0), "项目一", 1000, "问诊"])
    明细.append([external_customer_id, "Excel测试客户", phone, transaction_id, datetime(2026, 8, 26, 10, 0), "项目二", 500, "问诊"])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def real_abc_headers_workbook() -> bytes:
    """A memory-only workbook using the actual ABC 流水/分类/明细 headers."""
    workbook = Workbook()
    流水 = workbook.active
    流水.title = "流水"
    流水.append(["收费时间", "诊号", "患者姓名", "患者档案号", "性别", "年龄", "手机号", "支付方式", "实收", "科室", "结算险种", "结算类型", "医疗类别", "基金支付总额", "非基金支付总额"])
    流水.append([datetime(2026, 8, 31, 10, 0), "000001", "真实ABC客户", "000123", "女", 38, "+86 13800138000", "ABC支付(微信)", 1000, "中医科", "职工医保", "医保结算", "门诊", 700, 300])
    分类 = workbook.create_sheet("分类")
    分类.append(["收费时间", "患者姓名", "患者档案号", "费用分类", "实收"])
    分类.append([datetime(2026, 8, 31, 10, 0), "真实ABC客户", "000123", "治疗", 1000])
    明细 = workbook.create_sheet("明细")
    明细.append(["收费时间", "诊号", "患者姓名", "患者档案号", "手机号", "一级分类", "二级分类", "药品/项目名称", "药品/项目编码", "规格", "厂家", "数量", "单位", "单价", "实收", "项目备注"])
    明细.append([datetime(2026, 8, 31, 10, 0), "000001", "真实ABC客户", "000123", "+86 13800138000", "治疗", "理疗", "中医调理", "ITEM001", "一次", "ABC厂家", 1, "次", 1000, 1000, "真实明细"])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def zero_paid_deduction_workbook(external_customer_id: str, phone: str, transaction_id: str) -> bytes:
    workbook = Workbook()
    流水 = workbook.active
    流水.title = "流水"
    流水.append(["患者档案号", "患者姓名", "手机号", "收费单号", "收费时间", "诊号", "支付方式", "原价", "应收", "实收", "抵扣金额", "核销金额", "优惠金额", "项目名称", "数量", "单价", "明细金额"])
    流水.append([external_customer_id, "抵扣客户", phone, transaction_id, datetime(2026, 8, 31, 9, 0), "V-001", "现金", 258, 0, 0, -258, 0, 0, "三伏灸", 1, 258, 258])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def balance_transaction_workbook(external_customer_id: str, phone: str) -> bytes:
    workbook = Workbook()
    流水 = workbook.active
    流水.title = "流水"
    流水.append(["收费时间", "诊号", "患者姓名", "患者档案号", "手机号", "开单来源", "类型", "收费单号", "支付方式", "原价", "应收", "实收"])
    流水.append([datetime(2026, 8, 31, 10, 0), "V-001", "余额客户", external_customer_id, phone, "会员充值", "收费", "BAL-1", "ABC支付(微信)", 1000, 1000, 1000])
    流水.append([datetime(2026, 8, 31, 10, 5), "V-002", "余额客户", external_customer_id, phone, "会员充值", "退费", "BAL-2", "ABC支付(微信)", -200, -200, -200])
    流水.append([datetime(2026, 8, 31, 11, 0), "V-003", "余额客户", external_customer_id, phone, "治疗", "收费", "ORD-1", "ABC支付(微信)", 300, 300, 300])
    明细 = workbook.create_sheet("明细")
    明细.append(["收费时间", "诊号", "患者姓名", "患者档案号", "手机号", "一级分类", "药品/项目名称", "收费单号", "数量", "单位", "单价", "明细金额"])
    明细.append([datetime(2026, 8, 31, 10, 0), "V-001", "余额客户", external_customer_id, phone, "会员充值", "会员充值费", "BAL-1", 1, "次", 1000, 1000])
    明细.append([datetime(2026, 8, 31, 10, 5), "V-002", "余额客户", external_customer_id, phone, "会员充值", "会员充值退款", "BAL-2", 1, "次", -200, -200])
    明细.append([datetime(2026, 8, 31, 11, 0), "V-003", "余额客户", external_customer_id, phone, "治疗", "理疗项目", "ORD-1", 1, "次", 300, 300])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


class ExcelImportTest(unittest.TestCase):
    def setUp(self):
        self.db = SessionLocal()
        self.transaction = self.db.begin()
        self.admin = self.db.scalar(select(User).where(User.username == "admin01"))
        self.reporter = self.db.scalar(select(User).where(User.username == "sales01"))
        self.assignee = self.db.scalar(select(User).where(User.username == "sales02"))
        self.assertIsNotNone(self.admin)
        self.assertIsNotNone(self.reporter)
        self.assertIsNotNone(self.assignee)

    def tearDown(self):
        self.transaction.rollback()
        self.db.close()

    def _customer(self, phone: str) -> Customer:
        customer = Customer(
            name="内部客户原名",
            phone=phone,
            reported_by_id=self.reporter.id,
            assigned_to_id=self.assignee.id,
            source_category="online",
            lead_status="approved",
            status="active",
            pool_match_status="pending",
        )
        self.db.add(customer)
        self.db.flush()
        return customer

    def test_parse_and_confirm_matches_by_phone_and_keeps_internal_fields(self):
        suffix = uuid4().hex[:8]
        customer = self._customer(f"138{suffix}000")
        data = workbook_bytes(f"ABC-CUS-{suffix}", customer.phone, f"ABC-ORD-{suffix}")

        parsed = parse_excel(data, "abc.xlsx")
        self.assertEqual(parsed.sheet_names, ["流水", "明细"])
        self.assertEqual(len(parsed.transactions), 1)
        transaction = parsed.transactions[0]
        self.assertEqual(transaction.total_amount, Decimal("1500"))
        self.assertEqual(transaction.insurance_amount, Decimal("700"))
        self.assertEqual(transaction.self_pay_amount, Decimal("800"))
        self.assertEqual(len(transaction.items), 3)

        batch = preview_import(self.db, "abc.xlsx", data, self.admin.id)
        self.assertEqual(batch.matched_customers, 1)
        self.assertEqual(batch.unmatched_customers, 0)
        self.assertEqual((batch.preview_data or {})["evaluations"][0]["match_status"], "matched_auto")
        self.assertEqual((batch.preview_data or {})["evaluations"][0]["match_key"], "phone")
        confirm_import(self.db, batch, self.admin.id)
        self.db.flush()

        order = self.db.scalar(select(Order).where(Order.external_transaction_id == f"ABC-ORD-{suffix}"))
        self.assertIsNotNone(order)
        self.assertEqual(order.customer_id, customer.id)
        self.assertEqual(order.total_amount, Decimal("1500.00"))
        self.assertEqual(order.insurance_amount, Decimal("700.00"))
        self.assertEqual(order.self_pay_amount, Decimal("800.00"))
        self.assertEqual(self.db.scalar(select(func.count(OrderItem.id)).where(OrderItem.order_id == order.id)), 3)
        history = self.db.scalar(select(CustomerPoolMatchHistory).where(CustomerPoolMatchHistory.customer_id == customer.id, CustomerPoolMatchHistory.match_type == "auto"))
        self.assertIsNotNone(history)
        self.assertEqual(history.note, "abc_excel:phone")
        self.db.refresh(customer)
        self.assertEqual(customer.external_customer_id, f"ABC-CUS-{suffix}")
        self.assertEqual(customer.name, "Excel测试客户")
        self.assertEqual(customer.reported_by_id, self.reporter.id)
        self.assertEqual(customer.assigned_to_id, self.assignee.id)

    def test_matched_import_resolves_stale_pending_record(self):
        """A formal match also closes an old pending queue row."""
        suffix = uuid4().hex[:8]
        external_id = f"ABC-STALE-{suffix}"
        phone = f"138{suffix}222"
        data = workbook_bytes(external_id, phone, f"ABC-STALE-ORD-{suffix}")

        # The preview was created while the customer did not exist.  This
        # simulates an older import path that left an audit row pending even
        # though the customer became available before confirmation.
        batch = preview_import(self.db, "stale-pending.xlsx", data, self.admin.id)
        customer = self._customer(phone)
        transaction = parse_excel(data, "stale-pending.xlsx").transactions[0]
        stale = ExcelImportUnmatchedRecord(
            batch_id=batch.id,
            external_customer_id=external_id,
            name=transaction.name,
            phone=phone,
            transaction_fingerprint=transaction.transaction_fingerprint,
            raw_data={"normalized": transaction.to_dict()},
            status="pending",
            customer_id=customer.id,
        )
        self.db.add(stale)
        self.db.flush()

        # The queue API must hide a legacy pending row that already points to
        # a formal customer, even before the import confirmation repairs it.
        visible = unmatched_records(status=None, batch_id=None, db=self.db, operator=self.admin)
        self.assertNotIn(stale.id, {item["id"] for item in visible})

        confirm_import(self.db, batch, self.admin.id)
        self.db.flush()
        self.db.refresh(stale)

        self.assertEqual(stale.status, "bound")
        self.assertEqual(stale.customer_id, customer.id)
        self.assertIsNotNone(stale.resolved_at)
        self.assertEqual(
            self.db.scalar(select(Order.customer_id).where(Order.external_transaction_id == f"ABC-STALE-ORD-{suffix}")),
            customer.id,
        )

    def test_repeated_upload_skips_existing_order_and_unknown_customer_stays_unmatched(self):
        suffix = uuid4().hex[:8]
        customer = self._customer(f"139{suffix}000")
        data = workbook_bytes(f"ABC-CUS-{suffix}", customer.phone, f"ABC-ORD-{suffix}")
        first = preview_import(self.db, "abc.xlsx", data, self.admin.id)
        confirm_import(self.db, first, self.admin.id)
        self.db.flush()
        second = preview_import(self.db, "abc-again.xlsx", data, self.admin.id)
        confirm_import(self.db, second, self.admin.id)
        self.assertEqual(second.new_transactions, 0)
        self.assertEqual(second.skipped_transactions, 1)
        self.assertEqual(self.db.scalar(select(func.count(Order.id)).where(Order.external_transaction_id == f"ABC-ORD-{suffix}")), 1)

        unknown = workbook_bytes(f"ABC-UNKNOWN-{suffix}", f"137{suffix}000", f"ABC-UNKNOWN-ORD-{suffix}")
        unknown_batch = preview_import(self.db, "unknown.xlsx", unknown, self.admin.id)
        self.assertEqual(unknown_batch.unmatched_customers, 1)
        confirm_import(self.db, unknown_batch, self.admin.id)
        self.db.flush()
        unknown_placeholder = self.db.scalar(select(Customer).where(Customer.external_customer_id == f"ABC-UNKNOWN-{suffix}"))
        self.assertIsNotNone(unknown_placeholder)
        self.assertEqual(unknown_placeholder.pool_match_status, "pending")
        self.assertEqual(self.db.scalar(select(func.count(Order.id)).where(Order.customer_id == unknown_placeholder.id)), 1)
        self.assertEqual(unknown_batch.status, "partial")

        # Re-uploading the same unresolved transaction must not create a
        # second review row; the queue remains one customer/one transaction.
        repeated_unknown = preview_import(self.db, "unknown-again.xlsx", unknown, self.admin.id)
        confirm_import(self.db, repeated_unknown, self.admin.id)
        self.db.flush()
        self.assertEqual(
            self.db.scalar(
                select(func.count(ExcelImportUnmatchedRecord.id)).where(
                    ExcelImportUnmatchedRecord.transaction_fingerprint == parse_excel(unknown, "unknown.xlsx").transactions[0].transaction_fingerprint,
                )
            ),
            1,
        )
        grouped_unknown = [item for item in unmatched_records(status=None, batch_id=None, db=self.db, operator=self.admin) if item["external_customer_id"] == f"ABC-UNKNOWN-{suffix}"]
        self.assertEqual(len(grouped_unknown), 1)
        self.assertEqual(grouped_unknown[0]["transaction_count"], 1)

    def test_same_external_customer_pending_transactions_are_one_group(self):
        suffix = uuid4().hex[:8]
        external_id = f"ABC-GROUP-{suffix}"
        phone = f"137{suffix}000"
        data = balance_transaction_workbook(external_id, phone)
        batch = preview_import(self.db, "grouped.xlsx", data, self.admin.id)
        confirm_import(self.db, batch, self.admin.id)
        self.db.flush()

        groups = [item for item in unmatched_records(status=None, batch_id=None, db=self.db, operator=self.admin) if item["external_customer_id"] == external_id]
        self.assertEqual(len(groups), 1)
        self.assertEqual(groups[0]["transaction_count"], 3)
        self.assertEqual(groups[0]["total_amount"], "1100.00")
        self.assertEqual(len(groups[0]["transactions"]), 3)
        self.assertEqual(
            self.db.scalar(
                select(func.count(Order.id)).where(
                    Order.external_customer_id == external_id,
                )
            ),
            3,
        )

    def test_bound_customer_appends_new_transaction_without_reimporting_old_one(self):
        suffix = uuid4().hex[:8]
        external_id = f"ABC-APPEND-{suffix}"
        phone = f"137{suffix}111"
        customer = self._customer(phone)
        first_data = workbook_bytes(external_id, phone, f"ABC-APPEND-1-{suffix}")
        first = preview_import(self.db, "append-1.xlsx", first_data, self.admin.id)
        confirm_import(self.db, first, self.admin.id)
        self.db.flush()

        second_data = workbook_bytes(external_id, phone, f"ABC-APPEND-2-{suffix}")
        second = preview_import(self.db, "append-2.xlsx", second_data, self.admin.id)
        self.assertEqual(second.matched_customers, 1)
        self.assertEqual(second.unmatched_customers, 0)
        confirm_import(self.db, second, self.admin.id)
        self.db.flush()

        self.assertEqual(second.new_transactions, 1)
        self.assertEqual(
            self.db.scalar(select(func.count(Order.id)).where(Order.external_customer_id == external_id)),
            2,
        )
        self.assertEqual(
            self.db.scalar(select(func.count(Order.id)).where(Order.customer_id == customer.id)),
            2,
        )

        repeated = preview_import(self.db, "append-2-again.xlsx", second_data, self.admin.id)
        confirm_import(self.db, repeated, self.admin.id)
        self.db.flush()
        self.assertEqual(repeated.new_transactions, 0)
        self.assertEqual(repeated.skipped_transactions, 1)
        self.assertEqual(
            self.db.scalar(select(func.count(Order.id)).where(Order.external_customer_id == external_id)),
            2,
        )

    def test_excel_first_then_salesperson_lead_links_history_automatically(self):
        suffix = uuid4().hex[:8]
        external_id = f"ABC-FIRST-{suffix}"
        phone = f"135{suffix}000"
        batch = preview_import(self.db, "abc-first.xlsx", workbook_bytes(external_id, phone, f"ABC-FIRST-ORD-{suffix}"), self.admin.id)
        self.assertEqual(batch.matched_customers, 0)
        self.assertEqual(batch.unmatched_customers, 1)
        confirm_import(self.db, batch, self.admin.id)
        self.db.flush()

        placeholder = self.db.scalar(select(Customer).where(Customer.external_customer_id == external_id))
        self.assertIsNotNone(placeholder)
        order = self.db.scalar(select(Order).where(Order.external_order_id == f"ABC-FIRST-ORD-{suffix}"))
        self.assertIsNotNone(order)
        self.assertEqual(order.customer_id, placeholder.id)
        unmatched = self.db.scalar(select(ExcelImportUnmatchedRecord).where(ExcelImportUnmatchedRecord.batch_id == batch.id))
        self.assertEqual(unmatched.status, "pending")

        lead = Customer(
            name="后报备客户",
            phone=phone,
            reported_by_id=self.reporter.id,
            lead_status="pending",
            status="active",
            member_level="普通",
        )
        self.db.add(lead)
        self.db.flush()
        link_pending_external_customer(self.db, lead)
        self.db.flush()

        self.assertEqual(lead.external_customer_id, external_id)
        self.assertEqual(lead.reported_by_id, self.reporter.id)
        self.assertIsNone(lead.assigned_to_id)
        self.assertEqual(self.db.scalar(select(Order.customer_id).where(Order.id == order.id)), lead.id)
        self.db.refresh(unmatched)
        self.assertEqual(unmatched.status, "bound")
        self.assertEqual(unmatched.customer_id, lead.id)
        self.assertIsNotNone(self.db.scalar(select(CustomerPoolMatchHistory).where(CustomerPoolMatchHistory.customer_id == lead.id, CustomerPoolMatchHistory.match_type == "auto_lead_phone")))
        self.assertIsNone(self.db.scalar(select(Customer).where(Customer.id == placeholder.id)))

    def test_payment_mismatch_is_preserved_and_logged(self):
        suffix = uuid4().hex[:8]
        customer = self._customer(f"136{suffix}000")
        data = workbook_bytes(f"ABC-MISMATCH-{suffix}", customer.phone, f"ABC-MISMATCH-ORD-{suffix}", mismatch=True)
        batch = preview_import(self.db, "mismatch.xlsx", data, self.admin.id)
        confirm_import(self.db, batch, self.admin.id)
        self.db.flush()
        order = self.db.scalar(select(Order).where(Order.external_transaction_id == f"ABC-MISMATCH-ORD-{suffix}"))
        self.assertEqual(order.payment_split_status, "mismatch")
        self.assertEqual(order.insurance_amount + order.self_pay_amount, Decimal("900.00"))
        self.assertIsNotNone(self.db.scalar(select(IntegrationLog).where(IntegrationLog.event_type == "payment_amount_mismatch", IntegrationLog.external_id == f"ABC-MISMATCH-ORD-{suffix}")))

    def test_zero_paid_deduction_is_a_real_consumption_record(self):
        suffix = uuid4().hex[:8]
        customer = self._customer(f"138{suffix}111")
        batch = preview_import(self.db, "zero-deduction.xlsx", zero_paid_deduction_workbook(f"ABC-ZERO-{suffix}", customer.phone, f"ABC-ZERO-ORD-{suffix}"), self.admin.id)
        transaction = parse_excel(zero_paid_deduction_workbook(f"ABC-ZERO-{suffix}", customer.phone, f"ABC-ZERO-ORD-{suffix}"), "zero-deduction.xlsx").transactions[0]
        self.assertTrue(transaction.has_consumption_record)
        self.assertEqual(transaction.total_amount, Decimal("0"))
        self.assertEqual(transaction.paid_amount, Decimal("0"))
        self.assertEqual(transaction.original_amount, Decimal("258"))
        self.assertEqual(transaction.deduction_amount, Decimal("-258"))
        self.assertEqual(transaction.deduction_amount_raw, Decimal("-258"))
        confirm_import(self.db, batch, self.admin.id)
        self.db.flush()
        order = self.db.scalar(select(Order).where(Order.external_transaction_id == f"ABC-ZERO-ORD-{suffix}"))
        self.assertIsNotNone(order)
        self.assertTrue(order.has_consumption_record)
        self.assertEqual(order.total_amount, Decimal("0.00"))
        self.assertEqual(order.paid_amount, Decimal("0.00"))
        self.assertEqual(order.deduction_amount, Decimal("-258.00"))
        self.assertEqual(order.deduction_amount_raw, Decimal("-258.00"))
        self.db.refresh(customer)
        self.assertEqual(customer.external_order_count, 1)
        self.assertEqual(customer.external_total_spend, Decimal("258.00"))
        self.db.refresh(order)
        summary = calculate_customer_consumption_summary([order])
        self.assertEqual(summary.total_consumption, Decimal("258.00"))
        self.assertEqual(summary.actual_paid_amount, Decimal("0.00"))
        self.assertEqual(summary.effective_revenue, Decimal("258.00"))
        self.assertEqual(summary.insurance_amount, Decimal("0.00"))
        self.assertEqual(summary.self_pay_amount, Decimal("258.00"))
        self.assertEqual(summary.consumption_count, 1)
        self.assertEqual(summary.average_order_value, Decimal("258.00"))

    def test_real_abc_three_sheet_headers_enter_preview(self):
        parsed = parse_excel(real_abc_headers_workbook(), "收费明细2026-08-31_2026-08-31(2).xlsx")
        self.assertEqual(parsed.sheet_names, ["流水", "分类", "明细"])
        self.assertEqual(parsed.patient_count, 1)
        self.assertEqual(len(parsed.transactions), 1)
        self.assertEqual(parsed.transactions[0].external_customer_id, "000123")
        self.assertEqual(parsed.transactions[0].visit_no, "000001")
        self.assertEqual(parsed.transactions[0].payment_type, "self_pay")
        self.assertEqual(parsed.transactions[0].payment_type_raw, "ABC支付(微信)")
        self.assertEqual(len(parsed.transactions[0].items), 1)
        self.assertEqual({item["sheet_name"] for item in parsed.sheet_stats}, {"流水", "分类", "明细"})
        self.assertTrue(all(item["recognized"] for item in parsed.sheet_stats))
        self.assertTrue(any(item["header"] == "药品/项目名称" for item in next(item for item in parsed.sheet_stats if item["sheet_name"] == "明细")["recognized_fields"]))

        batch = preview_import(self.db, "收费明细2026-08-31_2026-08-31(2).xlsx", real_abc_headers_workbook(), self.admin.id)
        self.assertEqual(batch.status, "preview")
        self.assertEqual((batch.preview_data or {}).get("sheet_stats")[0]["sheet_name"], "流水")

    def test_member_recharge_and_refund_are_marked_and_excluded_from_consumption(self):
        suffix = uuid4().hex[:8]
        customer = self._customer(f"138{suffix}222")
        data = balance_transaction_workbook(f"ABC-BAL-{suffix}", customer.phone)
        parsed = parse_excel(data, "balance.xlsx")
        self.assertEqual(len(parsed.transactions), 3)
        kinds = {item.external_transaction_id: item.transaction_kind for item in parsed.transactions}
        self.assertEqual(kinds["BAL-1"], "balance_transaction")
        self.assertEqual(kinds["BAL-2"], "balance_transaction")
        self.assertEqual(kinds["ORD-1"], "consumption")

        batch = preview_import(self.db, "balance.xlsx", data, self.admin.id)
        confirm_import(self.db, batch, self.admin.id)
        self.db.flush()
        orders = self.db.scalars(select(Order).where(Order.customer_id == customer.id).order_by(Order.id)).all()
        self.assertEqual(len(orders), 3)
        self.assertEqual(sum(order.transaction_kind == "balance_transaction" for order in orders), 2)
        summary = calculate_customer_consumption_summary(orders)
        self.assertEqual(summary.total_consumption, Decimal("300"))
        self.assertEqual(summary.effective_revenue, Decimal("300"))
        self.assertEqual(summary.consumption_count, 1)


if __name__ == "__main__":
    unittest.main()
